import csv
import io
import json
import re
from pathlib import Path
from typing import Any, Iterator, Dict, Callable, Optional
import pandas as pd
from sqlalchemy import text
from sqlalchemy.orm import Session
import openpyxl

STANDARD_CRM_39_COLUMNS = [
    "mx_Campus",
    "CreatedOn",
    "ProspectID",
    "mx_FastTrackId",
    "FirstName",
    "Source",
    "SourceMedium",
    "Origin",
    "mx_State",
    "mx_State_New",
    "mx_Gender_New",
    "mx_Category_Backend",
    "mx_City",
    "mx_City_New",
    "mx_First_Allocation_Date_and_Time",
    "mx_First_Call_Disposition_Date",
    "mx_First_Call_Disposition",
    "mx_First_Call_Sub_Disposition",
    "mx_Latest_Follow_Up_Date",
    "mx_Call_Disposition",
    "mx_Call_Sub_Disposition",
    "mx_Last_Call_Date_and_Time",
    "mx_CUCET_First_Payment_Date",
    "mx_CUCET_Booked_Slot_Date",
    "mx_CUCET_Exam_Status",
    "mx_CUCET_Percentile",
    "mx_CUCET_Score",
    "mx_CUCET_Attempt_Counter",
    "mx_Slot_Date_CUCET",
    "mx_Scholarship_Perentage",
    "mx_Admission_done",
    "mx_AdmissionDate",
    "mx_Account_No",
    "ProspectStage",
    "OwnerIdName",
    "mx_Total_Call_Attempt",
    "mx_Refund_Initiated_On",
    "mx_Refund_Status",
    "Program Code",
]

STANDARD_CRM_41_COLUMNS = [
    "mx_Campus",
    "CreatedOn",
    "ProspectID",
    "mx_FastTrackId",
    "FirstName",
    "Source",
    "SourceMedium",
    "Origin",
    "mx_State",
    "mx_State_New",
    "mx_Gender_New",
    "mx_Category_Backend",
    "mx_City",
    "mx_City_New",
    "mx_State_dup",
    "mx_State_New_dup",
    "mx_First_Allocation_Date_and_Time",
    "mx_First_Call_Disposition_Date",
    "mx_First_Call_Disposition",
    "mx_First_Call_Sub_Disposition",
    "mx_Latest_Follow_Up_Date",
    "mx_Call_Disposition",
    "mx_Call_Sub_Disposition",
    "mx_Last_Call_Date_and_Time",
    "mx_CUCET_First_Payment_Date",
    "mx_CUCET_Booked_Slot_Date",
    "mx_CUCET_Exam_Status",
    "mx_CUCET_Percentile",
    "mx_CUCET_Score",
    "mx_CUCET_Attempt_Counter",
    "mx_Slot_Date_CUCET",
    "mx_Scholarship_Perentage",
    "mx_Admission_done",
    "mx_AdmissionDate",
    "mx_Account_No",
    "ProspectStage",
    "OwnerIdName",
    "mx_Total_Call_Attempt",
    "mx_Refund_Initiated_On",
    "mx_Refund_Status",
    "Program Code",
]


def is_headerless_crm_row(row: list) -> bool:
    """Detect if a row from a CSV is actually a CRM data record instead of a header."""
    if not row or len(row) < 30:
        return False
    col1 = str(row[1]).strip() if len(row) > 1 and row[1] else ""
    has_date_col1 = bool(re.match(r"^\d{4}-\d{2}-\d{2}", col1))
    col2 = str(row[2]).strip() if len(row) > 2 and row[2] else ""
    has_uuid_col2 = bool(re.match(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$", col2))
    col0 = str(row[0]).strip().lstrip("\ufeff") if len(row) > 0 and row[0] else ""
    is_not_header_col0 = not col0.lower().startswith("mx_") and not ("campus" in col0.lower() and "_" in col0.lower())
    return (has_date_col1 or has_uuid_col2) and is_not_header_col0


def clean_value_for_json(value: Any) -> Any:
    """Convert pandas/numpy values into JSON-compatible python types."""
    if pd.isna(value):
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            pass
    return value


def stream_file_records(file_path: str, chunk_size: int = 20000) -> Iterator[list[Dict[str, Any]]]:
    """
    Stream file contents in chunks to maintain low memory usage.
    Supports CSV, XLSX, XLS, and XLSB.
    """
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".csv":
        # Check if first line is headerless CRM data
        read_kwargs: Dict[str, Any] = {"chunksize": chunk_size}
        try:
            with open(file_path, "r", encoding="utf-8-sig", errors="replace") as f_check:
                first_row = next(csv.reader(f_check), None)
                if first_row and is_headerless_crm_row(first_row):
                    if len(first_row) == 39:
                        read_kwargs["names"] = STANDARD_CRM_39_COLUMNS
                        read_kwargs["header"] = None
                    elif len(first_row) == 41:
                        read_kwargs["names"] = STANDARD_CRM_41_COLUMNS
                        read_kwargs["header"] = None
        except Exception:
            pass

        for chunk in pd.read_csv(file_path, **read_kwargs):
            chunk_records = []
            for _, row in chunk.iterrows():
                rec = {str(col): clean_value_for_json(row[col]) for col in chunk.columns}
                chunk_records.append(rec)
            yield chunk_records

    elif extension == ".xlsx":
        wb = openpyxl.load_workbook(file_path, read_only=True)
        current_chunk = []
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            header = None
            rows = ws.iter_rows(values_only=True)
            for r in rows:
                if not header:
                    header = [str(cell) if cell is not None else f"col_{i}" for i, cell in enumerate(r)]
                    continue
                rec = {header[i]: clean_value_for_json(r[i]) if i < len(r) else None for i in range(len(header))}
                rec["sheet_name"] = sheetname
                current_chunk.append(rec)
                if len(current_chunk) >= chunk_size:
                    yield current_chunk
                    current_chunk = []
        if current_chunk:
            yield current_chunk
        wb.close()

    elif extension in [".xls", ".xlsb"]:
        engine = "pyxlsb" if extension == ".xlsb" else "xlrd"
        xl = pd.ExcelFile(file_path, engine=engine)
        current_chunk = []
        for sheetname in xl.sheet_names:
            df = pd.read_excel(xl, sheet_name=sheetname)
            for _, row in df.iterrows():
                rec = {str(col): clean_value_for_json(row[col]) for col in df.columns}
                rec["sheet_name"] = sheetname
                current_chunk.append(rec)
                if len(current_chunk) >= chunk_size:
                    yield current_chunk
                    current_chunk = []
        if current_chunk:
            yield current_chunk

    else:
        raise ValueError(f"Unsupported file format: {extension}")


class CsvToStagingStream(io.TextIOBase):
    """
    Memory-efficient streaming buffer for PostgreSQL COPY STDIN.
    Reads CSV file line-by-line formatted as tab-separated TSV:
    dataset_id \t row_number \t raw_data_json \t cleaning_status \n
    """
    def __init__(
        self,
        file_path: str,
        dataset_id: str,
        chunk_size: int = 50000,
        progress_callback: Optional[Callable[[int, int], None]] = None,
        total_rows: int = 0,
    ):
        self.file_obj = open(file_path, "r", encoding="utf-8-sig", errors="replace")
        self.reader = csv.reader(self.file_obj)
        self.dataset_id = str(dataset_id)
        self.chunk_size = chunk_size
        self.progress_callback = progress_callback
        self.total_rows = total_rows
        self.row_counter = 1
        self.pending_first_row: Optional[list] = None

        header_row = next(self.reader, None)
        if header_row:
            if is_headerless_crm_row(header_row):
                if len(header_row) == 39:
                    self.header = list(STANDARD_CRM_39_COLUMNS)
                elif len(header_row) == 41:
                    self.header = list(STANDARD_CRM_41_COLUMNS)
                else:
                    self.header = [
                        STANDARD_CRM_39_COLUMNS[i] if i < len(STANDARD_CRM_39_COLUMNS) else f"col_{i}"
                        for i in range(len(header_row))
                    ]
                self.pending_first_row = header_row
            else:
                self.header = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(header_row)]
        else:
            self.header = []
        self.buffer = io.StringIO()

    def read(self, size=-1):
        buf_val = self.buffer.getvalue()
        pos = self.buffer.tell()
        if pos < len(buf_val):
            chunk = buf_val[pos:]
            self.buffer.seek(0, io.SEEK_END)
            return chunk

        self.buffer = io.StringIO()
        writer = csv.writer(
            self.buffer,
            delimiter="\t",
            quotechar='"',
            quoting=csv.QUOTE_ALL,
            lineterminator="\n",
        )
        count = 0

        # Emit the first record if line 1 was detected as data (headerless CRM CSV)
        if self.pending_first_row:
            rec = {
                self.header[i]: self.pending_first_row[i] if i < len(self.pending_first_row) else None
                for i in range(len(self.header))
            }
            rec_json = json.dumps(rec, ensure_ascii=False)
            writer.writerow([self.dataset_id, self.row_counter, rec_json, "pending"])
            self.row_counter += 1
            count += 1
            self.pending_first_row = None
        for row in self.reader:
            if not row:
                continue
            rec = {
                self.header[i]: row[i] if i < len(row) else None
                for i in range(len(self.header))
            }
            rec_json = json.dumps(rec, ensure_ascii=False)
            writer.writerow([self.dataset_id, self.row_counter, rec_json, "pending"])
            self.row_counter += 1
            count += 1
            if count >= self.chunk_size:
                break

        if count > 0 and self.progress_callback:
            try:
                self.progress_callback(self.row_counter - 1, self.total_rows)
            except Exception:
                pass

        if count == 0:
            return ""

        self.buffer.seek(0)
        return self.buffer.read()

    def close(self):
        if hasattr(self, "file_obj") and self.file_obj:
            self.file_obj.close()


def load_to_staging(
    db: Session,
    dataset_id,
    file_path: str,
    batch_size: int = 20000,
    progress_callback: Optional[Callable[[int, int], None]] = None,
    total_rows: int = 0,
) -> int:
    """
    Load data from file into staging.records using PostgreSQL COPY bulk streaming.
    """
    from app.database.connection import engine

    path = Path(file_path)
    extension = path.suffix.lower()
    dataset_id_str = str(dataset_id)

    raw_conn = engine.raw_connection()
    try:
        cursor = raw_conn.cursor()
        copy_sql = "COPY staging.records (dataset_id, row_number, raw_data, cleaning_status) FROM STDIN WITH (FORMAT csv, DELIMITER E'\\t', QUOTE '\"', ESCAPE '\"')"

        if extension == ".csv":
            stream_buf = CsvToStagingStream(
                file_path,
                dataset_id_str,
                chunk_size=50000,
                progress_callback=progress_callback,
                total_rows=total_rows,
            )
            cursor.copy_expert(copy_sql, stream_buf)
            staged = stream_buf.row_counter - 1
            stream_buf.close()
            raw_conn.commit()

            if progress_callback:
                try:
                    progress_callback(staged, total_rows or staged)
                except Exception:
                    pass
            return staged

        else:
            row_counter = 1
            for chunk in stream_file_records(file_path, chunk_size=batch_size):
                buf = io.StringIO()
                writer = csv.writer(
                    buf,
                    delimiter="\t",
                    quotechar='"',
                    quoting=csv.QUOTE_ALL,
                    doublequote=True,
                    lineterminator="\n",
                )
                for record in chunk:
                    rec_json = json.dumps(record, ensure_ascii=False)
                    writer.writerow([dataset_id_str, row_counter, rec_json, "pending"])
                    row_counter += 1
                buf.seek(0)
                cursor.copy_expert(copy_sql, buf)
                buf.close()
                if progress_callback:
                    try:
                        progress_callback(row_counter - 1, total_rows)
                    except Exception:
                        pass
            raw_conn.commit()
            staged = row_counter - 1
            if progress_callback:
                try:
                    progress_callback(staged, total_rows or staged)
                except Exception:
                    pass
            return staged
    finally:
        raw_conn.close()